from django.shortcuts import get_object_or_404, redirect, render
from .models import *
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.db.models import Q, Count, Sum, query_utils
from .forms import GroupeForm, MembreGroupeForm, OrdreMembreForm, PaiementForm
from .models import Groupe, Paiement, Tour
from datetime import timedelta
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from django.http import HttpResponse


# Create your views here.
@login_required(login_url="login")
def dashboard_admin(request):

    if request.user.role != "ADMIN":
        return redirect("dashboard_membre")

    # =========================
    # DONNÉES DU TABLEAU DE BORD
    # =========================

    groupes = Groupe.objects.filter(
        admin=request.user
    )

    membres = MembreGroupe.objects.filter(
        groupe__admin=request.user
    )

    paiements = Paiement.objects.filter(
        membre__groupe__admin=request.user
    )

    tours = Tour.objects.filter(
        groupe__admin=request.user
    )

    # =========================
    # ACTIVITÉS RÉCENTES
    # =========================

    activites_recentes = []

    # =========================
    # NOUVEAUX MEMBRES
    # =========================

    membres_recents = MembreGroupe.objects.filter(
        groupe__admin=request.user
    ).order_by("-date_inscription")[:5]

    for membre in membres_recents:

        activites_recentes.append({
            "type": "MEMBRE",
            "titre": "Nouveau membre inscrit",
            "description": f"{membre.nom_affiche} a rejoint {membre.groupe.nom}",
            "date": membre.date_inscription.date(),
        })

    # =========================
    # COTISATIONS REÇUES
    # =========================

    paiements_recents = Paiement.objects.filter(
        membre__groupe__admin=request.user,
        statut="PAYE"
    ).order_by("-date_paiement")[:5]

    for paiement in paiements_recents:

        activites_recentes.append({
            "type": "COTISATION",
            "titre": "Cotisation reçue",
            "description": f"{paiement.membre.nom_affiche} - {paiement.montant:,.0f} FCFA",
            "date": paiement.date_paiement,
        })

    # =========================
    # TOURS TERMINÉS
    # =========================

    tours_termines = Tour.objects.filter(
        groupe__admin=request.user,
        statut="PAYE"
    ).order_by("-date_tour")[:5]

    for tour in tours_termines:

        activites_recentes.append({
            "type": "TOUR",
            "titre": "Tour terminé",
            "description": f"{tour.groupe.nom} - Tour de {tour.membre.nom_affiche}",
            "date": tour.date_tour,
        })

    # =========================
    # PAIEMENTS EN RETARD
    # =========================

    paiements_en_retard = Paiement.objects.filter(
        membre__groupe__admin=request.user,
        statut="NON_PAYE"
    ).order_by("-date_paiement")[:5]

    for paiement in paiements_en_retard:

        activites_recentes.append({
            "type": "RETARD",
            "titre": "Paiement en retard",
            "description": f"{paiement.membre.nom_affiche} - {paiement.montant:,.0f} FCFA",
            "date": paiement.date_paiement,
        })

    # =========================
    # TRIER LES ACTIVITÉS
    # =========================

    activites_recentes = sorted(
        activites_recentes,
        key=lambda activite: activite["date"],
        reverse=True
    )[:5]

    # =========================
    # CONTEXT
    # =========================

    context = {
        "total_membres": membres.count(),
        "total_groupes": groupes.count(),
        "total_paiements": paiements.count(),
        "total_tours": tours.count(),

        "activites_recentes": activites_recentes,
    }

    return render(
        request,
        "pages/dashboard_admin.html",
        context
    ) 

# VIEW EXPORT EXCEL MODAL
@login_required(login_url="login")
def export_excel(request):

    # Vérifier que l'utilisateur est administrateur
    if request.user.role != "ADMIN":
        return redirect("dashboard_membre")

    # =========================
    # RÉCUPÉRER LES DONNÉES
    # =========================

    groupes = Groupe.objects.filter(
        admin=request.user
    )

    membres = MembreGroupe.objects.filter(
        groupe__admin=request.user
    ).select_related(
        "groupe",
        "utilisateur"
    )

    paiements = Paiement.objects.filter(
        membre__groupe__admin=request.user
    ).select_related(
        "membre",
        "membre__groupe",
        "tour"
    )

    tours = Tour.objects.filter(
        groupe__admin=request.user
    ).select_related(
        "groupe",
        "membre"
    )

    # =========================
    # CRÉER LE FICHIER EXCEL
    # =========================

    workbook = Workbook()

    # =========================
    # FEUILLE GROUPES
    # =========================

    feuille_groupes = workbook.active
    feuille_groupes.title = "Groupes"

    entetes_groupes = [
        "Nom du groupe",
        "Montant cotisation",
        "Fréquence",
        "Date début",
        "Statut",
    ]

    feuille_groupes.append(entetes_groupes)

    for groupe in groupes:

        feuille_groupes.append([
            groupe.nom,
            float(groupe.montant_cotisation),
            groupe.get_frequence_display(),
            groupe.date_debut,
            groupe.get_statut_display(),
        ])

    # =========================
    # FEUILLE MEMBRES
    # =========================

    feuille_membres = workbook.create_sheet("Membres")

    entetes_membres = [
        "Membre",
        "Téléphone",
        "Groupe",
        "Ordre de réception",
        "Date d'inscription",
    ]

    feuille_membres.append(entetes_membres)

    for membre in membres:

        # Retirer la timezone du DateTimeField
        date_inscription = membre.date_inscription

        if date_inscription:
            date_inscription = timezone.make_naive(
                date_inscription
            )

        feuille_membres.append([
            membre.nom_affiche,
            membre.telephone,
            membre.groupe.nom if membre.groupe else "",
            membre.ordre_reception,
            date_inscription,
        ])

    # =========================
    # FEUILLE TOURS
    # =========================

    feuille_tours = workbook.create_sheet("Tours")

    entetes_tours = [
        "Groupe",
        "Membre",
        "Date du tour",
        "Statut",
    ]

    feuille_tours.append(entetes_tours)

    for tour in tours:

        feuille_tours.append([
            tour.groupe.nom,
            tour.membre.nom_affiche,
            tour.date_tour,
            tour.get_statut_display(),
        ])

    # =========================
    # FEUILLE PAIEMENTS
    # =========================

    feuille_paiements = workbook.create_sheet("Paiements")

    entetes_paiements = [
        "Membre",
        "Groupe",
        "Tour",
        "Montant",
        "Date paiement",
        "Statut",
    ]

    feuille_paiements.append(entetes_paiements)

    for paiement in paiements:

        feuille_paiements.append([
            paiement.membre.nom_affiche,
            paiement.membre.groupe.nom,
            paiement.tour.date_tour,
            float(paiement.montant),
            paiement.date_paiement,
            paiement.get_statut_display(),
        ])

    # =========================
    # STYLE DES EN-TÊTES
    # =========================

    feuilles = [
        feuille_groupes,
        feuille_membres,
        feuille_tours,
        feuille_paiements,
    ]

    for feuille in feuilles:

        # Style des en-têtes
        for cellule in feuille[1]:

            cellule.font = Font(
                bold=True,
                color="FFFFFF"
            )

            cellule.fill = PatternFill(
                fill_type="solid",
                fgColor="4F46E5"
            )

            cellule.alignment = Alignment(
                horizontal="center"
            )

        # Ajuster automatiquement la largeur
        for colonne in feuille.columns:

            longueur = 0

            lettre_colonne = colonne[0].column_letter

            for cellule in colonne:

                if cellule.value is not None:

                    longueur = max(
                        longueur,
                        len(str(cellule.value))
                    )

            feuille.column_dimensions[
                lettre_colonne
            ].width = longueur + 3

    # =========================
    # RÉPONSE HTTP
    # =========================

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="likelemba_export.xlsx"'
    )

    workbook.save(response)

    return response
@login_required(login_url="login")
def dashboard_membre(request):

    if request.user.role != 'MEMBRE':
        return redirect('dashboard_admin')

    return render(request, 'pages/dashboard_membre.html')

# Fonction pour afficher la liste de groupes
@login_required
def liste_groupes_view(request):

    if request.user.role != 'ADMIN':
        return redirect("dashboard_membre")
    
    groupes = Groupe.objects.filter(admin=request.user)
    return render(request, "groupes/liste_groupes.html", {'groupes': groupes})

#Fonction pour ajouter un groupe
@login_required(login_url="login")
def create_groupe(request):

    # Vérification ADMIN pour la sécurité 
    if request.user.role != 'ADMIN':
        return redirect('dashboard_membre')

    form = GroupeForm()

    if request.method == 'POST':
        form = GroupeForm(request.POST)

        if form.is_valid():

            groupe = form.save(commit=False)

            # Pour éviter que le groupe soit créé sans admin, on assigne l'admin avant de sauvegarder
            groupe.admin = request.user

            groupe.save()

            messages.success(request, "Groupe créé avec succès !")

            return redirect('liste_groupes')

        else:
            messages.error(request, "Veuillez corriger les erreurs dans le formulaire")

    return render(request, 'groupes/create_groupe.html', {
        'form': form
    })

# Fonction pour modifier un groupe
@login_required(login_url="login")
def update_groupe(request, id):

    groupe = get_object_or_404(Groupe, id=id, admin=request.user)

    form = GroupeForm(instance=groupe)

    if request.method == 'POST':
        form = GroupeForm(request.POST, instance=groupe)

        if form.is_valid():
            form.save()
            messages.success(request, "Groupe modifié avec succès !")
            return redirect('liste_groupes')
        else:
            messages.error(request, "Erreur lors de la modification !")

    return render(request, 'groupes/update_groupe.html', {
        'form': form
    })
# Fonction pour supprimer un groupe
@login_required(login_url="login")
def delete_groupe(request, id):

    groupe = get_object_or_404(Groupe, id=id, admin=request.user)

    if request.method == 'POST':
        groupe.delete()
        messages.success(request, "Groupe supprimé avec succès !")
        return redirect('liste_groupes')

    return redirect('liste_groupes')


@login_required(login_url='login')
def dashboard_financier_view(request, groupe_id):

    groupe = get_object_or_404(
        Groupe,
        id=groupe_id,
        admin=request.user
    )

    nombre_membres = MembreGroupe.objects.filter(
        groupe=groupe
    ).count()

    nombre_tours = Tour.objects.filter(
        groupe=groupe
    ).count()

    total_collecte = Paiement.objects.filter(
        tour__groupe=groupe
    ).aggregate(
        total=Sum('montant')
    )['total'] or 0

    total_attendu = (
        groupe.montant_cotisation *
        nombre_membres *
        nombre_tours
    )

    montant_restant = total_attendu - total_collecte

    taux_paiement = 0

    if total_attendu > 0:
        taux_paiement = round(
            (total_collecte / total_attendu) * 100,
            2
        )

    context = {
        'groupe': groupe,
        'nombre_membres': nombre_membres,
        'nombre_tours': nombre_tours,
        'total_collecte': total_collecte,
        'total_attendu': total_attendu,
        'montant_restant': montant_restant,
        'taux_paiement': taux_paiement,
    }

    return render(
        request,
        'dashboard/dashboard_financier.html',
        context
    )

# Fonction pour voir detail d'un groupe
@login_required(login_url="login")
def detail_groupe(request, id):
    groupe = get_object_or_404(Groupe, id=id, admin=request.user)
    
    # Calcul des statistiques
    total_membres = MembreGroupe.objects.filter(groupe=groupe).count()
    total_tours = Tour.objects.filter(groupe=groupe).count()
    total_paiements = Paiement.objects.filter(membre__groupe=groupe).count()
    
    return render(request, 'groupes/detail_groupe.html', {
        'groupe': groupe,
        'total_membres': total_membres,
        'total_tours': total_tours,
        'total_paiements': total_paiements,
    })

# Ajouter un membre dans un groupe
@login_required(login_url="login")
def ajouter_membre(request, groupe_id):
    groupe = get_object_or_404(Groupe, id=groupe_id, admin=request.user)

    if request.method == 'POST':
        form = MembreGroupeForm(request.POST, groupe=groupe)

        if form.is_valid():
            membre = form.save(commit=False)
            # Pour éviter que le membre soit crée sans groupe
            membre.groupe = groupe

            # éviter doublon
            if membre.utilisateur and MembreGroupe.objects.filter(utilisateur=membre.utilisateur, groupe=groupe).exists():
                messages.error(request, "Ce membre est déjà dans le groupe")
                return redirect('ajouter_membre', groupe_id=groupe.id)

            membre.save()
            messages.success(request, "Membre ajouté avec succès")
            return redirect('membres_groupe', groupe_id=groupe.id)
    else:
        form = MembreGroupeForm(groupe=groupe)

    return render(request, 'membres/ajouter_membre.html', {
        'form': form,
        'groupe': groupe
    })

@login_required(login_url="login")
def liste_membres_view(request):
    if request.user.role != 'ADMIN':
        return redirect('dashboard_membre')
    
    groupes = Groupe.objects.filter(admin=request.user).annotate(
        total_membres=Count('membregroupe')
    )

    return render(request, 'membres/liste_membres.html', {
        'groupes': groupes
    })

@login_required(login_url="login")
def membres_groupe_view(request, groupe_id):
    if request.user.role != 'ADMIN':
        return redirect('dashboard_membre')

    groupe = get_object_or_404(Groupe, id=groupe_id, admin=request.user)
    membres = MembreGroupe.objects.filter(groupe=groupe).order_by('ordre_reception', 'id')

    return render(request, 'membres/membres_groupe.html', {
        'groupe': groupe,
        'membres': membres
    })

@login_required(login_url="login")
def supprimer_membre(request, id):
    membre = get_object_or_404(
        MembreGroupe,
        id=id,
        groupe__admin=request.user  # 🔐 sécurité
    )

    if request.method == 'POST':
        groupe_id = membre.groupe.id
        membre.delete()
        messages.success(request, "Membre supprimé")
        return redirect('membres_groupe', groupe_id=groupe_id)

@login_required(login_url="login")
def modifier_ordre_membre(request, id):
    membre = get_object_or_404(
        MembreGroupe,
        id=id,
        groupe__admin=request.user  # 🔐 sécurité
    )

    if request.method == 'POST':
        form = OrdreMembreForm(request.POST, instance=membre, groupe=membre.groupe)

        if form.is_valid():
            form.save()
            messages.success(request, "Ordre modifié avec succès")
            return redirect('membres_groupe', groupe_id=membre.groupe.id)
    else:
        form = OrdreMembreForm(instance=membre, groupe=membre.groupe)

    return render(request, 'membres/modifier_ordre.html', {
        'form': form,
        'membre': membre
    })

@login_required(login_url="login")
def liste_paiements_groupes_view(request):
    if request.user.role != 'ADMIN':
        return redirect('dashboard_membre')

    groupes = Groupe.objects.filter(admin=request.user).annotate(
        total_membres=Count('membregroupe', distinct=True),
        total_paiements=Count('membregroupe__paiement', distinct=True),
        total_payes=Count('membregroupe__paiement', filter=Q(membregroupe__paiement__statut='PAYE'), distinct=True),
        total_attente=Count('membregroupe__paiement', filter=Q(membregroupe__paiement__statut='NON_PAYE'), distinct=True)
    )
    return render(request, 'paiements/liste_groupes_paiements.html', {
        'groupes': groupes
    })

@login_required(login_url="login")
def paiements_groupe_view(request, groupe_id):
    if request.user.role != 'ADMIN':
        return redirect('dashboard_membre')

    groupe = get_object_or_404(Groupe, id=groupe_id, admin=request.user)
    paiements = Paiement.objects.filter(
        membre__groupe=groupe
    ).select_related('membre', 'membre__utilisateur').order_by('-date_paiement', '-id')

    resume = paiements.aggregate(total_encaisse=Sum('montant'))
    total_encaisse = resume['total_encaisse'] or 0
    total_paiements = paiements.count()
    total_attendu = groupe.montant_cotisation * MembreGroupe.objects.filter(groupe=groupe).count()

    return render(request, 'paiements/paiements_groupe.html', {
        'groupe': groupe,
        'paiements': paiements,
        'total_encaisse': total_encaisse,
        'total_paiements': total_paiements,
        'total_attendu': total_attendu,
    })

@login_required(login_url="login")
def ajouter_paiement(request, groupe_id):
    if request.user.role != 'ADMIN':
        return redirect('dashboard_membre')

    groupe = get_object_or_404(Groupe, id=groupe_id, admin=request.user)

    if request.method == 'POST':
        form = PaiementForm(request.POST, groupe=groupe)

        if form.is_valid():

            paiement = form.save(commit=False)

            tour = Tour.objects.filter(
                membre=paiement.membre,
                statut='EN_ATTENTE'
            ).first()

            if not tour:
                messages.error(
                    request,
                    "Aucun tour en attente trouvé."
                )
                return redirect(
                    'ajouter_paiement',
                    groupe_id=groupe.id
                )

            paiement.tour = tour
            paiement.save()

            tour.statut = 'PAYE'
            tour.save()

            messages.success(
                request,
                "Paiement enregistré avec succès."
            )

            return redirect(
                'paiements_groupe',
                groupe_id=groupe.id
            )
    else:
        form = PaiementForm(groupe=groupe)

    return render(request, 'paiements/ajouter_paiement.html', {
        'form': form,
        'groupe': groupe
    })

@login_required(login_url="login")
def liste_tours_groupes_view(request):
    if request.user.role != 'ADMIN':
        return redirect('dashboard_membre')

    groupes = Groupe.objects.filter(admin=request.user).annotate(
        total_tours=Count('tour', distinct=True),
        total_membres=Count('membregroupe', distinct=True),
        total_payes=Count('tour', filter=Q(tour__statut='PAYE'), distinct=True),
        total_attente=Count('tour', filter=Q(tour__statut='EN_ATTENTE'), distinct=True)
    )

    # Calcul du nombre de tours terminés (payés) et en cours (en attente)
    for groupe in groupes:
        groupe.tours_termines = groupe.total_payes
        groupe.tours_en_cours = groupe.total_attente

    # Calcul des statistiques globales
    total_groupes = groupes.count()
    total_tours = sum(g.total_tours for g in groupes)
    total_payes = sum(g.total_payes for g in groupes)
    total_attente = sum(g.total_attente for g in groupes)

    return render(request, 'tours/liste_groupes_tours.html', {
        'groupes': groupes,
        'total_groupes': total_groupes,
        'total_tours': total_tours,
        'total_payes': total_payes,
        'total_attente': total_attente,
    })
    if request.user.role != 'ADMIN':
        return redirect('dashboard_membre')

    groupes = Groupe.objects.filter(admin=request.user).annotate(
        total_tours=Count('tour', distinct=True),
        total_membres=Count('membregroupe', distinct=True),
        total_payes=Count('tour', filter=Q(tour__statut='PAYE'), distinct=True),
        total_attente=Count('tour', filter=Q(tour__statut='EN_ATTENTE'), distinct=True)
    )

    # Calcul du nombre de tours terminés (payés) et en cours (en attente)
    for groupe in groupes:
        groupe.tours_termines = groupe.total_payes
        groupe.tours_en_cours = groupe.total_attente

    return render(request, 'tours/liste_groupes_tours.html', {
        'groupes': groupes
    })
def generer_tours(groupe):
    membres = MembreGroupe.objects.filter(groupe=groupe).order_by('ordre_reception')

    date = groupe.date_debut

    for membre in membres:
        Tour.objects.create(
            groupe=groupe,
            membre=membre,
            date_tour=date
        )

        # avancer selon fréquence
        if groupe.frequence == 'MENSUEL':
            date = date + timedelta(days=30)
        else:
            date = date + timedelta(days=7)

@login_required(login_url="login")
def lancer_tours(request, groupe_id):
    groupe = get_object_or_404(Groupe, id=groupe_id, admin=request.user)

    if Tour.objects.filter(groupe=groupe).exists():
        messages.warning(request, "Les tours existent déjà")
    else:
        generer_tours(groupe)
        messages.success(request, "Tours générés avec succès")

    return redirect('liste_tours', groupe_id=groupe.id)

@login_required(login_url="login")
def liste_tours_view(request, groupe_id):
    groupe = get_object_or_404(Groupe, id=groupe_id, admin=request.user)

    # Récupérer les tours du groupe
    tours = Tour.objects.filter(groupe=groupe).select_related('membre').order_by('date_tour')

    # Compter les statistiques
    total_tours = tours.count()
    tours_payes = tours.filter(statut='PAYE').count()
    tours_attente = tours.filter(statut='EN_ATTENTE').count()
    
    # Nombre de membres dans le groupe
    total_membres = MembreGroupe.objects.filter(groupe=groupe).count()

    # Ajouter les propriétés au groupe pour le template
    groupe.total_membres = total_membres
    groupe.total_tours = total_tours
    groupe.tours_payes = tours_payes
    groupe.tours_attente = tours_attente
    groupe.taux_completion = int((tours_payes / total_tours * 100)) if total_tours > 0 else 0

    return render(request, 'tours/liste_tours.html', {
        'groupe': groupe,
        'tours': tours,
    })



# Vues pour les utilisateurs membres
@login_required
def groupes_membre_view(request):

    if request.user.role == 'ADMIN':
        return redirect('dashboard_admin')

    # Récupérer tous les groupes auxquels l'utilisateur participe
    groupes = MembreGroupe.objects.filter(utilisateur=request.user).select_related('groupe')
    
    context = {
        'groupes': groupes
    }
    return render(request, 'groupes/membres/liste_groupes_membre.html', context)

# Vues pour voir les paiements des membres 
@login_required(login_url='login')
def paiements_membre_view(request):
    if request.user.role != 'MEMBRE':
        return redirect('dashboard_admin')

    paiements = Paiement.objects.filter(
        membre__utilisateur=request.user
    ).select_related('membre', 'membre__groupe').order_by('-date_paiement', '-id')

    total_encaisse = paiements.aggregate(total=Sum('montant'))['total'] or 0
    nombre_paiements = paiements.count()

    context = {
        'paiements': paiements,
        'total_encaisse': total_encaisse,
        'nombre_paiements': nombre_paiements,
    }
    return render(request, 'paiements/liste_paiements_membre.html', context)